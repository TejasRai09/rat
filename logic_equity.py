"""
logic_equity.py
──────────────────────────────────────────────────────────────────────────
Fills the “Fund Ranking_Equity” sheet of Estimation_V5.xlsx with
performance / risk data pulled from Kotak Report AS ON 01-Jul-2025.xlsx.

Blocks covered
  • Largecap, Large & Mid, Midcap, Smallcap, Focused
  • BAF  (Balanced Advantage)
  • Multi Asset
  • Sectoral-Thematic

Metrics filled (columns E → L)
  1-Yr Ret, 3-Yr Ret, 5-Yr Ret, 10-Yr Ret, Sharpe, Alpha (Jenson), Beta, SD
"""

from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
import sys

# ───────────────────────── FILE DEFAULTS ────────────────────────────────
EST_FILE   = Path("Estimation_V5.xlsx")
KOTAK_FILE = Path("Kotak Report AS ON 01-Jul-2025.xlsx")
EST_SHEET  = "Fund Ranking_Equity"

# ────────────────── COLUMN-LETTER → INDEX HELPER ────────────────────────
def col_idx(col: str) -> int:
    """Excel column letters → 0-based index (A=0)."""
    idx = 0
    for ch in col.upper():
        idx = idx * 26 + (ord(ch) - 64)
    return idx - 1

# ───────────── DEFAULT COLUMN POSITIONS (most Kotak sheets) ─────────────
RET_DEFAULT  = {            # N, P, Q, R
    "1yr": col_idx("N"),
    "3yr": col_idx("P"),
    "5yr": col_idx("Q"),
    "10yr": col_idx("R"),
}
RISK_DEFAULT = {            # BB, BC, BD, BE  (SD, Beta, Sharpe, Jenson)
    "SD":      col_idx("BB"),
    "Beta":    col_idx("BC"),
    "Sharpe":  col_idx("BD"),
    "Alpha":   col_idx("BE"),   # Jenson
}

# ───────────── SHEET-SPECIFIC OVERRIDES YOU REQUESTED ───────────────────
RET_OVERRIDE = {
    # Balanced Advantage Fund → S, U, V, W
    "BAF": {
        "1yr": col_idx("S"),
        "3yr": col_idx("U"),
        "5yr": col_idx("V"),
        "10yr": col_idx("W"),
    },
    # Multi Asset → V, X, Y, Z
    "Multi Asset": {
        "1yr": col_idx("V"),
        "3yr": col_idx("X"),
        "5yr": col_idx("Y"),
        "10yr": col_idx("Z"),
    },
    # Sectoral-Thematic stays at default (N,P,Q,R) – include anyway for clarity
    "Sectoral-Thematic": RET_DEFAULT,
}

RISK_OVERRIDE = {
    # BAF sheet → BG, BH, BI, BJ  (SD, Beta, Sharpe, Jenson)
    "BAF": {
        "SD":     col_idx("BG"),
        "Beta":   col_idx("BH"),
        "Sharpe": col_idx("BI"),
        "Alpha":  col_idx("BJ"),
    },
    # Multi Asset sheet → BA, BB, BC, BD
    "Multi Asset": {
        "SD":     col_idx("BA"),
        "Beta":   col_idx("BB"),
        "Sharpe": col_idx("BC"),
        "Alpha":  col_idx("BD"),
    },
    # Sectoral-Thematic sheet → BB, BC, BD, BE
    "Sectoral-Thematic": RISK_DEFAULT,
}

# Destination columns in Estimation (E → L)
EST_COLS = ["E", "F", "G", "H", "I", "J", "K", "L"]

# ──────────────────────── CATEGORY DEFINITIONS ──────────────────────────
#  (unchanged except new Sectoral-Thematic block)
CATEGORIES = [
    {"sheet": "Largecap", "schemes": [
        "Aditya Birla Sun Life Large Cap Fund - Reg - Growth",
        "HDFC Large Cap Fund - Growth",
        "ICICI Prudential Large Cap Fund - Growth",
        "Kotak Large Cap Fund - Reg - Growth",
        "Nippon India Large Cap Fund - Reg - Growth",
        "SBI Large Cap Fund - Growth",
        "UTI Large Cap Fund - Growth",
    ]},
    {"sheet": "Large & Mid", "schemes": [
        "SBI Large & Midcap Fund - Growth",
        "Kotak Large & Midcap Fund - Reg - Growth",
        "HDFC Large and Mid Cap Fund - Growth",
        "ICICI Prudential Large & Mid Cap Fund - Growth",
        "Nippon India Vision Large & Mid Cap Fund - Reg - Growth",
        "Aditya Birla Sun Life Large & Mid Cap Fund - Growth",
        "UTI Large & Mid Cap Fund - Growth",
    ]},
    {"sheet": "Midcap", "schemes": [
        "HDFC Mid Cap Fund - Growth",
        "Kotak Midcap Fund - Reg - Growth",
        "Nippon India Growth Mid Cap Fund - Reg - Growth",
        "SBI Midcap Fund - Growth",
        "UTI Mid Cap Fund - Growth",
        "ICICI Prudential MidCap Fund - Growth",
        "Aditya Birla Sun Life Mid Cap Fund - Plan A - Growth",
    ]},
    {"sheet": "Smallcap", "schemes": [
        "Nippon India Small Cap Fund - Reg - Growth",
        "HDFC Small Cap Fund - Growth",
        "SBI Small Cap Fund - Growth",
        "Kotak Small Cap Fund - Reg - Growth",
        "ICICI Prudential Smallcap Fund - Growth",
        "Aditya Birla Sun Life Small Cap Fund - Growth",
        "UTI Small Cap Fund - Reg - Growth",
    ]},
    {"sheet": "Focused Fund", "schemes": [
        "SBI Focused Fund - Growth",
        "HDFC Focused Fund - Growth",
        "Axis Focused Fund - Growth",
        "ICICI Prudential Focused Equity Fund  - Ret - Growth",
        "Nippon India Focused Fund - Reg - Growth",
        "Kotak Focused Fund - Reg - Growth",
        "UTI Focused Fund - Reg - Growth",
    ]},
    {"sheet": "BAF", "schemes": [
        "HDFC Balanced Advantage Fund - Growth",
        "ICICI Prudential Balanced Advantage Fund - Reg - Growth",
        "SBI Balanced Advantage Fund - Reg - Growth",
        "Kotak Balanced Advantage Fund - Reg - Growth",
        "UTI Balanced Advantage Fund - Reg - Growth",
        "Nippon India Balanced Advantage Fund - Reg - Growth",
        "Aditya Birla Sun Life Balanced Advantage Fund - Growth",
    ]},
    {"sheet": "Multi Asset", "schemes": [
        "ICICI Prudential Multi-Asset Fund - Growth",
        "SBI Multi Asset Allocation Fund - Growth",
        "Kotak Multi Asset Allocation Fund -Reg - Growth",
        "Nippon India Multi Asset Allocation Fund - Reg - Growth",
        "UTI Multi Asset Allocation Fund - Growth",
        "HDFC Multi - Asset Fund - Growth",
        "Aditya Birla Sun Life Multi Asset Allocation Fund - Reg - Growth",
    ]},
    {"sheet": "Sectoral-Thematic", "schemes": [
        # Aditya Birla SL
            "Aditya Birla Sun Life PSU Equity Fund - Reg - Growth",
            "Aditya Birla Sun Life Banking and Financial Services Fund - Reg - Growth",
            "Aditya Birla Sun Life Quant Fund - Reg - Growth",
            "Aditya Birla Sun Life Business Cycle Fund - Reg - Growth",
            "Aditya Birla Sun Life Conglomerate Fund - Reg - Growth",
            "Aditya Birla Sun Life Transportation and Logistics Fund - Reg - Growth",
            "Aditya Birla Sun Life Manufacturing Equity Fund - Reg - Growth",
            "Aditya Birla Sun Life Pharma & Healthcare Fund - Reg - Growth",
            "Aditya Birla Sun Life Special Opportunities Fund - Reg - Growth",
            "Aditya Birla Sun Life ESG Integration Strategy Fund - Reg - Growth",

            # UTI
            "UTI Quant Fund - Reg - Growth",
            "UTI Innovation Fund - Reg - Growth",

            # ICICI Pru
            "ICICI Prudential Pharma Healthcare and Diagnostics Fund - Reg - Growth",
            "ICICI Prudential Equity Minimum Variance Fund - Reg - Growth",
            "ICICI Prudential Bharat Consumption Fund - Reg - Growth",
            "ICICI Prudential Transportation and Logistics Fund - Reg - Growth",
            "ICICI Prudential Commodities Fund - Reg - Growth",
            "ICICI Prudential Housing Opportunities Fund - Reg - Growth",
            "ICICI Prudential PSU Equity Fund - Reg - Growth",
            "ICICI Prudential Quality Fund - Reg - Growth",
            "ICICI Prudential MNC Fund - Reg - Growth",
            "ICICI Prudential ESG Exclusionary Strategy Fund - Reg - Growth",
            "ICICI Prudential Rural Opportunities Fund - Reg - Growth",
            "ICICI Prudential Quant Fund - Reg - Growth",
            "ICICI Prudential Active Momentum Fund - Reg - Growth",

            # HDFC
            "HDFC Banking & Financial Services Fund - Reg - Growth",
            "HDFC Business Cycle Fund - Reg - Growth",
            "HDFC Pharma and Healthcare Fund - Reg - Growth",
            "HDFC Technology Fund - Reg - Growth",
            "HDFC Transportation and Logistics Fund - Reg - Growth",
            "HDFC Housing Opportunities Fund - Reg - Growth",
            "HDFC Non-Cyclical Consumer Fund - Reg - Growth",
            "HDFC MNC Fund - Reg - Growth",
            "HDFC Innovation Fund - Reg - Growth",

            # SBI
            "SBI Automotive Opportunities Fund - Reg - Growth",
            "SBI Quant Fund - Reg - Growth",
            "SBI Equity Minimum Variance Fund - Reg - Growth",

            # Nippon India
            "Nippon India Innovation Fund - Reg - Growth",
            "Nippon India Consumption Fund - Reg - Growth",
            "Nippon India US Equity Opportunities Fund - Reg - Growth",
            "Nippon India Taiwan Equity Fund - Reg - Growth",
            "Nippon India Japan Equity Fund - Reg  - Growth",
            "Nippon India Active Momentum Fund - Reg - Growth",
            "Nippon India Quant Fund - Reg - Growth",

            # Kotak
            "Kotak Pioneer Fund - Reg - Growth",
            "Kotak Business Cycle Fund - Reg - Growth",
            "Kotak Manufacture in India Fund - Reg - Growth",
            "Kotak MNC Fund - Reg - Growth",
            "Kotak Special Opportunities Fund - Reg - Growth",
            "Kotak Consumption Fund - Reg - Growth",
            "Kotak Banking & Financial Services Fund - Reg - Growth",
            "Kotak ESG Exclusionary Strategy Fund - Reg - Growth",
            "Kotak Quant Fund - Reg - Growth",
            "Kotak Technology Fund - Reg - Growth",
            "Kotak Healthcare Fund - Reg - Growth",
            "Kotak Transportation & Logistics Fund - Reg - Growth",
            "Kotak Energy Opportunities Fund - Reg - Growth",
    ]},
]

# ────────────────────────── HELPERS ─────────────────────────────────────
def _clean(x):
    """ '--', 'NA', blanks → 0.0 ; else float """
    if pd.isna(x):
        return 0.0
    s = str(x).strip()
    if s.lower() in {"--", "-", "na", "n.a", "n/a"}:
        return 0.0
    try:
        return float(s.replace(",", ""))
    except ValueError:
        print(f"[WARN] cannot convert '{s}', using 0")
        return 0.0


def sheet_maps(sheet):
    """Return (ret_cols, risk_cols) dicts for the given sheet name."""
    return (
        RET_OVERRIDE.get(sheet, RET_DEFAULT),
        RISK_OVERRIDE.get(sheet, RISK_DEFAULT),
    )


def pull_metrics(sheet, schemes):
    """Return {scheme: [8 metrics]} from one Kotak worksheet."""
    df   = pd.read_excel(KOTAK_FILE, sheet_name=sheet, header=None)
    names = df[0].astype(str).str.strip()

    rets, risks = sheet_maps(sheet)
    out = {}
    for sc in schemes:
        mask = names.eq(sc)
        if not mask.any():
            print(f"[WARN] '{sc}' not found in '{sheet}'")
            continue
        r = mask.idxmax()
        raw = [
            df.iat[r, rets["1yr"]],
            df.iat[r, rets["3yr"]],
            df.iat[r, rets["5yr"]],
            df.iat[r, rets["10yr"]],
            df.iat[r, risks["Sharpe"]],
            df.iat[r, risks["Alpha"]],
            df.iat[r, risks["Beta"]],
            df.iat[r, risks["SD"]],
        ]
        out[sc] = [_clean(v) for v in raw]
    return out


def write_metrics(all_data, est_path):
    wb = load_workbook(est_path)
    ws = wb[EST_SHEET]

    # quick lookup: fund name → row
    name_to_row = {
        (ws.cell(r, 2).value or "").strip(): r
        for r in range(1, ws.max_row + 1)
    }

    for block in all_data.values():
        for scheme, metrics in block.items():
            row = name_to_row.get(scheme)
            if not row:
                print(f"[WARN] '{scheme}' missing in Estimation sheet")
                continue
            for col_letter, val in zip(EST_COLS, metrics):
                ws[f"{col_letter}{row}"].value = val

    wb.save(est_path)
    print("✓ Fund Ranking_Equity updated")


def update_estimation(kotak_path, est_path):
    global KOTAK_FILE, EST_FILE
    KOTAK_FILE = Path(kotak_path)
    EST_FILE   = Path(est_path)

    combined = {}
    for blk in CATEGORIES:
        combined[blk["sheet"]] = pull_metrics(blk["sheet"], blk["schemes"])
    write_metrics(combined, EST_FILE)


# ─────────────────────────── CLI ENTRY ──────────────────────────────────
if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python logic_equity.py <Kotak.xlsx> <Estimation.xlsx>")
        sys.exit(1)
    update_estimation(sys.argv[1], sys.argv[2])
