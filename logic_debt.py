import pandas as pd
from openpyxl import load_workbook
import shutil, os
import sys

# 17 metrics to extract for every block
METRICS = [
    "YTM", "Modified Duration", "Macaulay Duration",
    "3 Months", "6 Months", "9 Months",
    "1 Year", "2 Years", "3 Years", "5 Years", "10 Years",
    "Sovereign", "AAA/A1+", "AA+", "AA/AA-", "Below AA-", "Cash and Fixed Deposits"
]

# Configuration for each duration block: (sheet name, schemes list, header row, first output row)
BLOCKS = [
    ("Ultra Short Duration", [
        "Aditya Birla Sun Life Savings Fund - Reg - Growth",
        "HDFC Ultra Short Term Fund - Reg - Growth",
        "SBI Magnum Ultra Short Duration Fund - Growth",
        "ICICI Prudential Ultra Short Term Fund - Growth",
        "Kotak Savings Fund - Reg - Growth",
        "Nippon India Ultra Short Duration Fund - Reg - Growth",
        "UTI Ultra Short Duration Fund - Growth"
    ], 6, 7),
    ("Low Duration", [
        "ICICI Prudential Savings Fund - Reg - Growth",
        "HDFC Low Duration Fund - Growth",
        "SBI Magnum Low Duration Fund - Growth",
        "Aditya Birla Sun Life Low Duration Fund - Reg - Growth",
        "Kotak Low Duration Fund - Std - Growth",
        "Nippon India Low Duration Fund - Reg - Growth",
        "UTI Low Duration Fund - Reg - Growth"
    ], 18, 19),
    ("Short Duration", [
        "ICICI Prudential Short Term Fund - Growth",
        "Kotak Bond Short Term Fund - Reg - Growth",
        "HDFC Short Term Debt Fund - Growth",
        "SBI Short Term Debt Fund - Growth",
        "Aditya Birla Sun Life Short Term Fund - Reg - Growth",
        "Nippon India Short Duration Fund - Reg - Growth",
        "UTI Short Duration Fund - Reg - Growth"
    ], 30, 31),
    ("Medium Duration", [
        "SBI Magnum Medium Duration Fund - Growth",
        "ICICI Prudential Medium Term Bond Fund - Growth",
        "HDFC Medium Term Debt Fund - Growth",
        "Aditya Birla Sun Life Medium Term Plan - Reg - Growth",
        "Kotak Medium Term Fund - Reg - Growth",
        "Nippon India Medium Duration Fund - Reg - Growth",
        "UTI Medium Duration Fund - Reg - Growth"
    ], 41, 42),
    ("Medium to Long Duration", [
        "ICICI Prudential Bond Fund - Growth",
        "Aditya Birla Sun Life Income Fund - Reg - Growth",
        "Kotak Bond Fund - Reg - Growth",
        "SBI Magnum Income Fund - Growth",
        "HDFC Income Fund - Growth",
        "Nippon India Medium to Long Duration Fund - Reg - G P - Growth",
        "UTI Medium to Long Duration Fund - Growth"
    ], 52, 53)
]

def flatten_kotak(kotak_path, sheet_name):
    print(f"[DEBUG] Reading Kotak sheet: {sheet_name}")
    raw = pd.read_excel(kotak_path, sheet_name=sheet_name, header=[4, 5], dtype=str)
    raw.columns = [f"{a or ''}{' ' if a and b else ''}{b or ''}".strip() for a, b in raw.columns]
    print(f"[DEBUG] Raw columns: {raw.columns.tolist()}")

    df = raw.rename(columns={raw.columns[0]: "Scheme"})
    df["Scheme"] = df["Scheme"].str.strip()
    df = df.applymap(lambda x: 0 if isinstance(x, str) and x.strip() == "--" else x)
    print(f"[DEBUG] Cleaned columns: {df.columns.tolist()}")
    return df

def process(est_path, kotak_path):
    print(f"[DEBUG] Starting process with Estimation: {est_path}, Kotak: {kotak_path}")
    base = os.path.dirname(est_path)
    output = os.path.join(base, "Estimation_V4_updated.xlsx")
    shutil.copy(est_path, output)
    print(f"[DEBUG] Copied estimation to output: {output}")

    wb = load_workbook(output)
    ws = wb['Fund Ranking_Debt']

    for sheet, schemes, header_row, first_row in BLOCKS:
        print(f"[DEBUG] Processing block: {sheet}")
        df_k = flatten_kotak(kotak_path, sheet)
        print(f"[DEBUG] df_k columns after flatten: {df_k.columns.tolist()}")

        # Build metric→column mapping
        col_map = {}
        for m in METRICS:
            matches = [c for c in df_k.columns if m in c]
            if not matches:
                print(f"[DEBUG] No column match for metric: {m}")
            exact = [c for c in matches if c.endswith(m)]
            col_map[m] = exact[0] if exact else (matches[0] if matches else None)
        print(f"[DEBUG] Column map: {col_map}")

        # Extract and round values
        data = {}
        for sc in schemes:
            if sheet == "Short Duration":
                subset = df_k[df_k["Scheme"].str.contains(sc, na=False)]
            else:
                subset = df_k[df_k["Scheme"] == sc]
            if subset.empty:
                print(f"[DEBUG] Warning: '{sc}' not found in {sheet}")
                continue
            row = subset.iloc[0]
            data[sc] = {}
            for m in METRICS:
                raw_val = row[col_map[m]] if col_map[m] else None
                try:
                    val = round(float(raw_val), 2)
                except:
                    val = raw_val
                data[sc][m] = val
        print(f"[DEBUG] Extracted data keys: {list(data.keys())}")

        # Locate Estimation columns
        est_cols = {}
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=header_row, column=col).value
            if val in METRICS:
                est_cols[val] = col
        print(f"[DEBUG] Estimation columns map: {est_cols}")

        # Write values into Estimation sheet
        for idx, sc in enumerate(schemes):
            row_idx = first_row + idx
            if sc in data:
                for m, col_idx in est_cols.items():
                    ws.cell(row=row_idx, column=col_idx, value=data[sc][m])
        print(f"[DEBUG] Finished writing block: {sheet}")

    wb.save(output)
    print(f"[DEBUG] Saved updated estimation file: {output}")
    return output

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python logic.py <Estimation_V4.xlsx> <Kotak_Report.xlsx>")
    else:
        process(sys.argv[1], sys.argv[2])
